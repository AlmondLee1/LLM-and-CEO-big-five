from __future__ import annotations

import argparse
import asyncio
import codecs
import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import tomllib
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit


BASE_DIR = Path(__file__).resolve().parents[1]
PROVIDERS = ("gpt", "deepseek")
DIMENSIONS = ("openness", "conscientiousness", "extraversion", "agreeableness", "neuroticism")
TARGET_SCORES = 100
PROMPT_TEMPLATE = """Please rate the {{dimension}} personality trait of {{ceo_name}}, CEO of {{company}}.

Rating range: 1 to 7 (1 is the lowest; 7 is the highest).

Requirements:
1. Output exactly one numeric score with two decimal places.
2. Do not output any explanation, analysis, or additional text.
3. The score must be between 1 and 7, inclusive.

Score:"""


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def emit(message, progress=None):
    if progress is None:
        print(message, flush=True)
    else:
        progress(message)


def target_scores(run):
    mode = run.get("mode", "research")
    target = positive_integer(run["valid_scores_per_ceo"], "valid_scores_per_ceo")
    if mode not in ("research", "pilot", "demo"):
        raise ValueError("mode must be research, pilot or demo")
    return target


def build_prompt(person, dimension):
    if dimension not in DIMENSIONS:
        raise ValueError("Unknown personality dimension")
    return re.sub(r"\{\{(company|ceo_name|dimension)\}\}",
                  lambda match: {"company": person["company"], "ceo_name": person["name"],
                                 "dimension": dimension}[match[1]], PROMPT_TEMPLATE)


def parse_score(text):
    value = text.strip()
    return float(value) if re.fullmatch(r"(?:[1-6]\.[0-9]{2}|7\.00)", value) else None


def csv_reader(text):
    for delimiter in (",", ";", "\t"):
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        columns = [column.strip().lstrip("\ufeff").casefold() for column in (reader.fieldnames or [])]
        if {"ceo_id", "name", "company"} <= set(columns):
            if len(columns) != len(set(columns)):
                raise ValueError("Input CSV column names must be unique")
            reader.fieldnames = columns
            return reader
    raise ValueError("Input CSV needs ceo_id, name and company columns")


def read_csv_text(path, encoding="auto"):
    raw = path.read_bytes()
    if not raw:
        raise ValueError("Input CSV is empty")

    def decode(selected):
        try:
            text = raw.decode(selected)
        except (UnicodeError, LookupError) as exc:
            raise ValueError("Cannot read this CSV with the selected encoding; choose another CSV encoding in Advanced") from exc
        if "\x00" in text:
            raise ValueError("The CSV contains unreadable characters; check its encoding or export it as a text CSV")
        return text, selected

    if encoding != "auto":
        return decode(encoding)
    for signature, selected in ((codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
                                (codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"),
                                (codecs.BOM_UTF16_BE, "utf-16")):
        if raw.startswith(signature):
            return decode(selected)
    if b"\x00" in raw:
        candidates = []
        for selected in ("utf-16-le", "utf-16-be"):
            try:
                text, selected = decode(selected)
                csv_reader(text)
                candidates.append((text, selected))
            except (ValueError, csv.Error):
                continue
        if len(candidates) == 1:
            return candidates[0]
        raise ValueError("Cannot identify this CSV encoding; choose its encoding in Advanced")
    try:
        return decode("utf-8-sig")
    except ValueError:
        pass
    candidates = []
    for selected in ("cp1252", "gb18030"):
        try:
            candidates.append(decode(selected))
        except ValueError:
            continue
    if len(candidates) == 1 or (candidates and all(item[0] == candidates[0][0] for item in candidates)):
        return candidates[0]
    if candidates:
        raise ValueError("CSV encoding is ambiguous; choose Windows-1252 or GB18030 / GBK in Advanced, or use --input-encoding")
    raise ValueError("Cannot identify this CSV encoding; select the encoding used when the file was saved")


def read_ceos(path, encoding="auto"):
    text, _ = read_csv_text(path, encoding)
    try:
        reader = csv_reader(text)
        people = []
        identifiers = set()
        for number, row in enumerate(reader, 2):
            if None in row or any(not isinstance(row.get(key), str) or not row[key].strip()
                                  for key in ("ceo_id", "name", "company")):
                raise ValueError(f"Invalid CEO record at CSV row {number}")
            person = {key: row[key].strip() for key in ("ceo_id", "name", "company")}
            if person["ceo_id"] in identifiers:
                raise ValueError(f"Duplicate CEO ID at CSV row {number}")
            identifiers.add(person["ceo_id"])
            people.append(person)
    except csv.Error as exc:
        raise ValueError("The CSV has invalid quoting or inconsistent delimiters") from exc
    if not people:
        raise ValueError("Input CSV contains no CEOs")
    return people


def positive_integer(value, label):
    if type(value) is not int or value < 1:
        raise ValueError(f"{label} must be a positive integer")
    return value


def finite_number(value, label, minimum=0):
    if type(value) not in (int, float) or not math.isfinite(value) or value < minimum:
        raise ValueError(f"Invalid {label}")
    return float(value)


def validate_url(value, label):
    parsed = urlsplit(value)
    if (parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password
            or parsed.query or parsed.fragment):
        raise ValueError(f"{label} must be an HTTPS URL without credentials or query parameters")
    return value


def load_config(path):
    config = tomllib.loads(path.read_text(encoding="utf-8-sig"))
    return validate_config(config)


def validate_config(config):
    run = config.get("run", {})
    providers = config.get("providers", {})
    if set(providers) != set(PROVIDERS):
        raise ValueError("Configure exactly the gpt and deepseek providers")
    target_scores(run)
    encoding = run.get("input_encoding", "auto")
    if not isinstance(encoding, str):
        raise ValueError("input_encoding must be auto or a text encoding name")
    if encoding != "auto":
        try:
            codecs.lookup(encoding)
        except LookupError as exc:
            raise ValueError("Unsupported CSV encoding") from exc
    if run.get("provider") not in (*PROVIDERS, "all"):
        raise ValueError("Invalid run.provider")
    if run.get("dimension") not in (*DIMENSIONS, "all"):
        raise ValueError("Invalid run.dimension")
    for key in ("input", "output_dir"):
        if not isinstance(run.get(key), str) or not run[key].strip():
            raise ValueError(f"Set run.{key}")
    positive_integer(run.get("max_tokens"), "max_tokens")
    finite_number(run.get("timeout_seconds"), "timeout_seconds", 0.001)
    finite_number(run.get("request_delay_seconds"), "request_delay_seconds")
    finite_number(run.get("retry_delay_seconds"), "retry_delay_seconds", 0.001)
    finite_number(run.get("max_retry_delay_seconds"), "max_retry_delay_seconds", run["retry_delay_seconds"])
    for provider, settings in providers.items():
        for key in ("model", "api_key_env"):
            if not isinstance(settings.get(key), str) or not settings[key].strip():
                raise ValueError(f"Set providers.{provider}.{key}")
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", settings["api_key_env"]):
            raise ValueError("api_key_env must name an environment variable")
        positive_integer(settings.get("concurrency"), f"{provider}.concurrency")
        positive_integer(settings.get("max_tokens", run["max_tokens"]), f"{provider}.max_tokens")
        adapter = settings.get("adapter")
        if adapter == "openai_compatible":
            validate_url(settings.get("base_url", ""), "base_url")
        elif adapter == "azure_openai" and provider == "gpt":
            for key in ("azure_endpoint_env", "api_version_env"):
                if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", settings.get(key, "")):
                    raise ValueError(f"Set providers.gpt.{key}")
        else:
            raise ValueError(f"Unsupported adapter for {provider}")
        if settings.get("temperature_mode") not in ("omit", "set"):
            raise ValueError("temperature_mode must be omit or set")
        if settings["temperature_mode"] == "set":
            value = finite_number(settings.get("temperature"), "temperature")
            if value > 2:
                raise ValueError("temperature must be between 0 and 2")
        if provider == "deepseek" and settings.get("thinking_mode", "provider_default") not in (
                "enabled", "disabled", "provider_default"):
            raise ValueError("Invalid thinking_mode")
    return config


def load_local_env(path, names):
    if not path.is_file():
        return
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        name, separator, value = line.strip().partition("=")
        name, value = name.strip(), value.strip()
        if not separator or name not in names:
            continue
        if len(value) >= 2 and value[0] in ("'", '"') and value[-1] == value[0]:
            value = value[1:-1]
        if value:
            os.environ.setdefault(name, value)


def environment_value(name, credentials=None):
    value = (credentials or {}).get(name, os.getenv(name, "")).strip()
    if not value or value.lower().startswith(("your_", "your-", "replace_", "<")):
        raise ValueError(f"Set the environment variable {name}")
    return value


def request_parameters(provider, settings, run):
    options = {"model": settings["model"],
               "max_completion_tokens" if provider == "gpt" else "max_tokens":
                   settings.get("max_tokens", run["max_tokens"])}
    if settings["temperature_mode"] == "set":
        options["temperature"] = float(settings["temperature"])
    if provider == "deepseek" and settings.get("thinking_mode", "provider_default") != "provider_default":
        options["extra_body"] = {"thinking": {"type": settings["thinking_mode"]}}
    return options


def make_api_call(provider, settings, run, credentials=None):
    from openai import AsyncAzureOpenAI, AsyncOpenAI

    options = {"api_key": environment_value(settings["api_key_env"], credentials),
               "timeout": float(run["timeout_seconds"]), "max_retries": 0}
    if settings["adapter"] == "azure_openai":
        client = AsyncAzureOpenAI(
            **options, azure_endpoint=validate_url(environment_value(settings["azure_endpoint_env"], credentials),
                                                 "azure_endpoint"),
            api_version=environment_value(settings["api_version_env"], credentials))
    else:
        client = AsyncOpenAI(**options, base_url=settings["base_url"])

    async def call(prompt):
        response = await client.chat.completions.create(
            **request_parameters(provider, settings, run), messages=[{"role": "user", "content": prompt}])
        message = response.choices[0].message if response.choices else None
        text = (message.content or getattr(message, "refusal", "") or "") if message else ""
        return {"text": text, "model": response.model or settings["model"],
                "request_id": response.id,
                "usage": response.usage.model_dump() if response.usage else None}

    call.aclose = client.close
    return call


def make_demo_call(provider, settings, run):
    requests = 0

    async def call(prompt):
        nonlocal requests
        requests += 1
        number = requests
        await asyncio.sleep(0.005)
        if number % 23 == 0:
            raise TimeoutError()
        score = (100 + (int(sha256(prompt.encode())[:8], 16) + number * 13) % 601) / 100
        return {"text": "Example invalid response" if number % 17 == 0 else f"{score:.2f}",
                "model": "synthetic-demo", "request_id": f"demo-{number}", "usage": None}

    async def close():
        return None

    call.aclose = close
    return call


def run_metadata(people_hash, provider, dimension, settings, run, credentials=None):
    connection = {"adapter": settings["adapter"]}
    if settings["adapter"] == "azure_openai":
        endpoint = (credentials or {}).get(settings["azure_endpoint_env"],
                                            os.getenv(settings["azure_endpoint_env"], "")).strip()
        version = (credentials or {}).get(settings["api_version_env"],
                                           os.getenv(settings["api_version_env"], "")).strip()
        connection.update(endpoint_sha256=sha256(endpoint.encode()) if endpoint else None,
                          api_version=version or None)
    else:
        connection["base_url"] = settings["base_url"]
    return {"schema_version": 1, "protocol": "english-valid-scores-v1",
            "input_sha256": people_hash, "provider": provider, "dimension": dimension,
            "input_encoding": run.get("input_encoding", "auto"),
            "valid_scores_per_ceo": target_scores(run),
            "mode": "synthetic_demo" if run.get("mode") == "demo" else run.get("mode", "research"),
            "prompt_template": PROMPT_TEMPLATE, "prompt_sha256": sha256(PROMPT_TEMPLATE.encode()),
            "connection": connection, "request_parameters": request_parameters(provider, settings, run)}


def check_manifest(folder, metadata, resume, check_only=False):
    manifest = folder / "run.json"
    journal = folder / "attempts.jsonl"
    if not manifest.exists():
        if folder.exists() and any(p.name != ".writer.lock" for p in folder.iterdir()):
            raise ValueError("Output files exist without run.json; choose a new output directory")
        return
    if not resume and not check_only:
        raise FileExistsError("Output exists; use --resume or choose a new output directory")
    saved = json.loads(manifest.read_text(encoding="utf-8"))
    if any(saved.get(key, "auto" if key == "input_encoding" else None) != value for key, value in metadata.items()):
        raise ValueError("Saved run has incompatible inputs or settings; choose a new output directory")
    if journal.exists() and not journal.is_file():
        raise ValueError("Invalid attempts journal")


def load_journal(folder, people, dimension, provider, model, recover=False, target=TARGET_SCORES):
    accepted, counts = {}, {}
    path = folder / "attempts.jsonl"
    if not path.exists():
        return accepted, counts
    known = {p["ceo_id"]: p for p in people}
    complete_bytes = 0
    incomplete = b""
    with path.open("rb") as handle:
        for number, line in enumerate(handle, 1):
            if not line.endswith(b"\n"):
                incomplete = line
                break
            try:
                record = json.loads(line)
                slot = (record["ceo_id"], record["iteration"])
                person = known.get(slot[0])
                score = parse_score(record["response_text"])
                if (person is None or type(slot[1]) is not int or not 1 <= slot[1] <= target
                        or record["provider"] != provider or record["dimension"] != dimension
                        or record["model_requested"] != model or record["ceo_name"] != person["name"]
                        or record["company"] != person["company"] or slot in accepted
                        or type(record["attempt"]) is not int
                        or record["attempt"] != counts.get(slot, 0) + 1
                        or record["status"] not in ("ok", "api_error", "invalid_response")
                        or (record["status"] == "ok" and (score is None or type(record["score"]) not in
                                                                        (int, float) or record["score"] != score))
                        or (record["status"] != "ok" and record["score"] is not None)
                        or (record["status"] == "invalid_response" and score is not None)):
                    raise ValueError("Inconsistent attempt")
            except (KeyError, TypeError, ValueError, AttributeError) as exc:
                raise ValueError(f"Invalid attempts journal at line {number}") from exc
            counts[slot] = record["attempt"]
            if record["status"] == "ok":
                accepted[slot] = record
            complete_bytes += len(line)
    if incomplete and recover:
        backup = folder / f"attempts.interrupted-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%f')}.bin"
        backup.write_bytes(incomplete)
        with path.open("r+b") as handle:
            handle.truncate(complete_bytes)
        print("Recovered an interrupted journal write; the partial line was backed up.", flush=True)
    return accepted, counts


@contextmanager
def output_lock(folder):
    folder.mkdir(parents=True, exist_ok=True)
    with (folder / ".writer.lock").open("a+b") as handle:
        if handle.seek(0, 2) == 0:
            handle.write(b"0")
            handle.flush()
        handle.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise FileExistsError("Another process is using this output directory") from exc
        try:
            yield
        finally:
            handle.seek(0)
            if os.name == "nt":
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def atomic_text(path, text):
    temporary = path.with_name(f".{path.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(text)
        handle.flush()
        os.fsync(handle.fileno())
    temporary.replace(path)


def atomic_csv(path, headers, rows):
    temporary = path.with_name(f".{path.name}.tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
        handle.flush()
        os.fsync(handle.fileno())
    temporary.replace(path)


def write_outputs(folder, people, accepted, counts, target=TARGET_SCORES, export_dir=None):
    from openpyxl import Workbook

    manifest = json.loads((folder / "run.json").read_text(encoding="utf-8"))
    mode = manifest.get("mode", "research")
    destination = folder if export_dir is None else Path(export_dir)
    destination.mkdir(parents=True, exist_ok=True)
    stem = "ratings" if export_dir is None else result_stem(
        manifest["request_parameters"]["model"], manifest["dimension"], target, mode)
    ordered = [accepted[(p["ceo_id"], i)] for p in people for i in range(1, target + 1)
               if (p["ceo_id"], i) in accepted]
    if any(r["status"] != "ok" or parse_score(r["response_text"]) != r["score"]
           or not 1 <= r["score"] <= 7 for r in ordered):
        raise ValueError("Only valid scores may be exported")
    complete = len(ordered) == len(people) * target
    suffix = "" if complete else ".partial"
    atomic_text(folder / "ratings.jsonl", "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in ordered))
    rating_headers = ["ceo_id", "ceo_name", "company", "provider", "model_requested", "model_returned",
                      "dimension", "iteration", "score", "mode"]
    atomic_csv(folder / f"ratings{suffix}.csv", rating_headers,
               [[r.get(key, mode if key == "mode" else "") if key != "score" else f"{r['score']:.2f}"
                 for key in rating_headers] for r in ordered])
    book = Workbook()
    ratings = book.active
    ratings.title = "Ratings"
    ratings.append(["Iteration"] + [f"{p['name']} [{p['ceo_id']}]" for p in people])
    for iteration in range(1, target + 1):
        ratings.append([iteration] + [accepted.get((p["ceo_id"], iteration), {}).get("score") for p in people])
    headers = ["ceo_id", "ceo_name", "company", "valid", "target", "completed_attempts",
               "failed_attempts", "mean_valid_score", "complete", "mode"]
    summary = book.create_sheet("Summary")
    summary.append(headers)
    rows = []
    means = []
    for person in people:
        slots = [(person["ceo_id"], i) for i in range(1, target + 1)]
        scores = [accepted[s]["score"] for s in slots if s in accepted]
        attempts = sum(counts.get(s, 0) for s in slots)
        mean = round(sum(scores) / len(scores), 4) if scores else None
        means.append(mean)
        rows.append([person["ceo_id"], person["name"], person["company"], len(scores), target,
                     attempts, attempts - len(scores), mean, len(scores) == target, mode])
        summary.append(rows[-1])
        for cell in summary[summary.max_row][:3]:
            cell.data_type = "s"
    ratings.append(["Mean valid score"] + means)
    for row in ratings.iter_rows(min_row=2, max_row=target + 1, min_col=2):
        for cell in row:
            cell.number_format = "0.00"
    for cell in ratings[target + 2][1:]:
        cell.number_format = "0.0000"
    for row in summary.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = "0.0000"
    ratings.freeze_panes = "B2"
    summary.freeze_panes = "A2"
    info = book.create_sheet("Run info")
    info.append(["Setting", "Value"])
    for key in ("mode", "provider", "dimension", "valid_scores_per_ceo", "input_sha256", "prompt_sha256"):
        info.append([key, manifest.get(key)])
    info.append(["model_requested", manifest["request_parameters"]["model"]])
    temporary = destination / f".{stem}.tmp.xlsx"
    try:
        book.save(temporary)
        temporary.replace(destination / f"{stem}.xlsx")
    finally:
        book.close()
    atomic_csv(folder / f"summary{suffix}.csv", headers,
               [[*row[:7], f"{row[7]:.4f}" if row[7] is not None else "", *row[8:]] for row in rows])
    if complete:
        for name in ("ratings.partial.csv", "summary.partial.csv"):
            (folder / name).unlink(missing_ok=True)


async def collect_scores(people, dimension, provider, model, call, run, concurrency, folder, accepted, counts, progress=None):
    target = target_scores(run)
    queue = asyncio.Queue()
    for iteration in range(1, target + 1):
        for person in people:
            if (person["ceo_id"], iteration) not in accepted:
                queue.put_nowait((person, iteration))
    total = len(people) * target
    emit(f"{provider} / {dimension}: {len(accepted)}/{total} valid scores", progress)
    with (folder / "attempts.jsonl").open("a", encoding="utf-8", newline="\n") as journal:
        async def worker():
            while True:
                person, iteration = await queue.get()
                try:
                    slot = (person["ceo_id"], iteration)
                    attempt = counts.get(slot, 0) + 1
                    response, error, http_status, fatal = {}, None, None, False
                    try:
                        async with asyncio.timeout(float(run["timeout_seconds"])):
                            response = await call(build_prompt(person, dimension))
                    except Exception as exc:
                        error = type(exc).__name__
                        http_status = getattr(exc, "status_code", None)
                        fatal = (type(http_status) is int and 400 <= http_status < 500
                                 and http_status not in (408, 409, 429)) or isinstance(exc, (TypeError, ValueError))
                    text = response.get("text", "")
                    score = parse_score(text) if error is None else None
                    status = "api_error" if error else ("ok" if score is not None else "invalid_response")
                    record = {"ceo_id": person["ceo_id"], "ceo_name": person["name"],
                              "mode": "synthetic_demo" if run.get("mode") == "demo" else run.get("mode", "research"),
                              "company": person["company"], "dimension": dimension, "provider": provider,
                              "model_requested": model, "model_returned": response.get("model"),
                              "iteration": iteration, "attempt": attempt, "score": score, "status": status,
                              "response_text": text, "error_type": error, "http_status": http_status,
                              "request_id": response.get("request_id"), "usage": response.get("usage"),
                              "completed_utc": utc_now()}
                    journal.write(json.dumps(record, ensure_ascii=False) + "\n")
                    journal.flush()
                    os.fsync(journal.fileno())
                    counts[slot] = attempt
                    if score is not None:
                        accepted[slot] = record
                        if len(accepted) % max(1, min(100, total // 100)) == 0 or len(accepted) == total:
                            emit(f"{provider} / {dimension}: {len(accepted)}/{total} valid scores", progress)
                    elif fatal:
                        raise ValueError(f"API request stopped ({error}, HTTP {http_status}); correct the configuration and resume")
                    else:
                        emit(f"Pending: {person['ceo_id']} / {dimension} / {iteration}; {status}; retrying", progress)
                    delay = float(run["request_delay_seconds"])
                    if score is None:
                        delay = max(delay, min(float(run["retry_delay_seconds"]) * 2 ** min(attempt - 1, 16),
                                               float(run["max_retry_delay_seconds"])))
                    if delay:
                        await asyncio.sleep(delay)
                    if score is None:
                        queue.put_nowait((person, iteration))
                finally:
                    queue.task_done()

        async with asyncio.TaskGroup() as group:
            workers = [group.create_task(worker()) for _ in range(min(concurrency, queue.qsize()))]
            await queue.join()
            for worker_task in workers:
                worker_task.cancel()
    if len(accepted) != total:
        raise ValueError("Collection ended before all valid scores were obtained")


async def run_collection(people, dimension, provider, settings, run, folder, metadata, resume,
                         call_factory=make_api_call, progress=None, export_dir=None):
    target = target_scores(run)
    with output_lock(folder):
        check_manifest(folder, metadata, resume)
        if not (folder / "run.json").exists():
            atomic_text(folder / "run.json", json.dumps({**metadata, "created_utc": utc_now()}, indent=2) + "\n")
        accepted, counts = load_journal(folder, people, dimension, provider, settings["model"], recover=True, target=target)
        call = None
        try:
            if len(accepted) < len(people) * target:
                call = call_factory(provider, settings, run)
                await collect_scores(people, dimension, provider, settings["model"], call, run,
                                     settings["concurrency"], folder, accepted, counts, progress)
        finally:
            try:
                if call is not None:
                    await call.aclose()
            finally:
                write_outputs(folder, people, accepted, counts, target, export_dir)
        result = folder if export_dir is None else Path(export_dir) / f"{result_stem(model=settings['model'], dimension=dimension, target=target, mode=run.get('mode', 'research'))}.xlsx"
        emit(f"Complete: {provider} / {dimension}; {len(people)} CEOs x {target} valid scores; {result}", progress)


def result_stem(model, dimension, target, mode="research"):
    model = re.sub(r"[^A-Za-z0-9._-]+", "_", model).strip("._-") or "model"
    prefix = "demo-" if mode in ("demo", "synthetic_demo") else ("pilot-" if mode == "pilot" else "")
    return f"{prefix}{model}-{dimension}-{target}"


def output_folder(root, provider, dimension, settings, target=TARGET_SCORES, mode="research"):
    return root / result_stem(settings["model"], dimension, target, mode)


def collection_root(root, input_path, input_hash, run):
    return root / "records"


def prepare_collections(config, input_path, root, providers, dimensions, resume=True,
                        check_only=False, credentials=None):
    validate_config(config)
    run = config["run"]
    target = target_scores(run)
    people = read_ceos(input_path, run.get("input_encoding", "auto"))
    people_hash = sha256(input_path.read_bytes())
    destination = collection_root(root, input_path, people_hash, run)
    tasks = []
    for provider in providers:
        settings = config["providers"][provider]
        if not check_only and run.get("mode") != "demo":
            environment_value(settings["api_key_env"], credentials)
            if settings["adapter"] == "azure_openai":
                validate_url(environment_value(settings["azure_endpoint_env"], credentials), "azure_endpoint")
                environment_value(settings["api_version_env"], credentials)
        for dimension in dimensions:
            folder = output_folder(destination, provider, dimension, settings, target, run.get("mode", "research"))
            metadata = run_metadata(people_hash, provider, dimension, settings, run, credentials)
            check_manifest(folder, metadata, resume, check_only=check_only)
            load_journal(folder, people, dimension, provider, settings["model"], target=target)
            tasks.append((provider, dimension, settings, folder, metadata))
    return people, tasks


async def main_async(args):
    config_path = args.config.resolve()
    config = load_config(config_path)
    run = config["run"]
    if args.mode is not None:
        run["mode"] = args.mode
        run["valid_scores_per_ceo"] = TARGET_SCORES if args.mode == "research" else 2
    if args.scores is not None:
        run["valid_scores_per_ceo"] = args.scores
    if args.input_encoding is not None:
        run["input_encoding"] = args.input_encoding
    target = target_scores(run)
    names = {s[key] for s in config["providers"].values()
             for key in ("api_key_env", "azure_endpoint_env", "api_version_env") if key in s}
    if run.get("mode") != "demo":
        load_local_env(config_path.parent / ".env", names)
    input_path = (args.input if args.input is not None else config_path.parent / run["input"]).resolve()
    root = (args.output_dir if args.output_dir is not None else config_path.parent / run["output_dir"]).resolve()
    provider_choice = args.provider or run["provider"]
    dimension_choice = args.dimension or run["dimension"]
    providers = PROVIDERS if provider_choice == "all" else (provider_choice,)
    dimensions = DIMENSIONS if dimension_choice == "all" else (dimension_choice,)
    for dependency in (("openpyxl",) if run.get("mode") == "demo" else ("openpyxl", "openai")):
        __import__(dependency)

    people, tasks = prepare_collections(config, input_path, root, providers, dimensions,
                                       resume=args.resume, check_only=args.check)
    print(f"Input: {len(people)} CEOs; target: {target} valid scores per CEO, trait and model; {len(tasks)} collections", flush=True)
    if args.check:
        print("Configuration and input checks passed. No API requests were made.", flush=True)
        return
    for provider, dimension, settings, folder, metadata in tasks:
        if run.get("mode") == "demo":
            run = {**run, "request_delay_seconds": 0, "retry_delay_seconds": 0.01,
                   "max_retry_delay_seconds": 0.01}
        await run_collection(people, dimension, provider, settings, run, folder, metadata, args.resume,
                             call_factory=make_demo_call if run.get("mode") == "demo" else make_api_call,
                             export_dir=root)


def error_leaves(exc):
    if isinstance(exc, BaseExceptionGroup):
        return [leaf for child in exc.exceptions for leaf in error_leaves(child)]
    return [exc]


def main():
    parser = argparse.ArgumentParser(description="Collect valid English-prompt Big Five ratings per CEO, trait and model.")
    parser.add_argument("--config", type=Path, default=BASE_DIR / "config" / "config.toml")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--provider", choices=(*PROVIDERS, "all"))
    parser.add_argument("--dimension", choices=(*DIMENSIONS, "all"))
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--mode", choices=("research", "pilot", "demo"))
    parser.add_argument("--scores", type=int)
    parser.add_argument("--input-encoding")
    args = parser.parse_args()
    try:
        asyncio.run(main_async(args))
    except KeyboardInterrupt:
        print("Stopped. Saved valid scores can be continued with --resume.", file=sys.stderr)
        return 130
    except Exception as exc:
        for leaf in error_leaves(exc):
            if isinstance(leaf, (ValueError, FileNotFoundError, FileExistsError, ImportError)):
                print(f"Error: {leaf}", file=sys.stderr)
            else:
                print(f"Error: {type(leaf).__name__}; check configuration, dependencies and local output files.", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
